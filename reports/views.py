import csv
from io import BytesIO
from urllib.parse import urlencode

from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Count, Q
from django.http import HttpResponse
from django.urls import reverse
from django.views.generic import TemplateView

from feedback.models import StakeholderFeedback
from monitoring.models import ProjectActivity
from users.models import User
from users.permissions import filter_workers_for_user, filter_workplaces_for_user
from workers.choices import DISABILITY_TYPE_CHOICES, SEVERITY_CHOICES
from workers.models import Worker, Workplace

from .forms import REPORT_CHOICES, ReportFilterForm


REPORT_LABELS = dict(REPORT_CHOICES)


class ReportIndexView(LoginRequiredMixin, TemplateView):
    template_name = "reports/index.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        form = ReportFilterForm(self.request.GET or {"report": "workers"})
        if not form.is_valid():
            form = ReportFilterForm({"report": "workers"})
            form.is_valid()

        report = build_report(self.request, form.cleaned_data, limit=200)
        querystring = report_querystring(form.cleaned_data)
        context.update(
            {
                "filter_form": form,
                "report": report,
                "export_csv_url": export_url("reports:workers-csv", querystring),
                "export_xlsx_url": export_url("reports:workers-xlsx", querystring),
                "export_pdf_url": export_url("reports:summary-pdf", querystring),
            }
        )
        return context


def export_url(name, querystring):
    url = reverse(name)
    return f"{url}?{querystring}" if querystring else url


def report_querystring(cleaned_data):
    params = {
        key: value
        for key, value in cleaned_data.items()
        if value not in ("", None)
    }
    params.setdefault("report", "workers")
    return urlencode(params)


def _worker_queryset(request):
    return filter_workers_for_user(
        Worker.objects.select_related("workplace", "disability_profile", "labour_rights", "created_by"),
        request.user,
    )


def _workplace_queryset(request):
    return filter_workplaces_for_user(
        Workplace.objects.annotate(worker_count=Count("workers", distinct=True)),
        request.user,
    )


def _feedback_queryset(request):
    queryset = StakeholderFeedback.objects.select_related("user")
    user = request.user
    if user.is_superuser or user.role in {User.ROLE_SYSTEM_ADMIN, User.ROLE_NATIONAL_MANAGER}:
        return queryset
    return queryset.filter(user=user)


def _apply_worker_filters(queryset, filters):
    q = filters.get("q")
    if q:
        queryset = queryset.filter(
            Q(unique_id__icontains=q)
            | Q(first_name__icontains=q)
            | Q(last_name__icontains=q)
            | Q(other_names__icontains=q)
            | Q(phone__icontains=q)
            | Q(workplace__employer_name__icontains=q)
            | Q(workplace__workplace_name__icontains=q)
            | Q(occupation__icontains=q)
        )
    if filters.get("province"):
        queryset = queryset.filter(province=filters["province"])
    if filters.get("district"):
        queryset = queryset.filter(district=filters["district"])
    if filters.get("gender"):
        queryset = queryset.filter(gender=filters["gender"])
    if filters.get("sector"):
        queryset = queryset.filter(sector=filters["sector"])
    if filters.get("employment_status"):
        queryset = queryset.filter(employment_status=filters["employment_status"])
    if filters.get("union"):
        queryset = queryset.filter(union_name=filters["union"])
    if filters.get("disability_type"):
        queryset = queryset.filter(disability_profile__disability_type=filters["disability_type"])
    if filters.get("verified") == "yes":
        queryset = queryset.filter(data_verified=True)
    if filters.get("verified") == "no":
        queryset = queryset.filter(data_verified=False)
    if filters.get("capture_source"):
        queryset = queryset.filter(capture_source=filters["capture_source"])
    if filters.get("capture_mode"):
        queryset = queryset.filter(capture_mode=filters["capture_mode"])
    return queryset


def _apply_workplace_filters(queryset, filters):
    q = filters.get("q")
    if q:
        queryset = queryset.filter(
            Q(employer_name__icontains=q)
            | Q(workplace_name__icontains=q)
            | Q(district__icontains=q)
            | Q(province__icontains=q)
            | Q(address__icontains=q)
        )
    if filters.get("province"):
        queryset = queryset.filter(province=filters["province"])
    if filters.get("district"):
        queryset = queryset.filter(district=filters["district"])
    if filters.get("sector"):
        queryset = queryset.filter(sector=filters["sector"])
    return queryset


def _apply_activity_filters(queryset, filters):
    q = filters.get("q")
    if q:
        queryset = queryset.filter(
            Q(title__icontains=q)
            | Q(partner__icontains=q)
            | Q(district__icontains=q)
            | Q(province__icontains=q)
            | Q(notes__icontains=q)
        )
    if filters.get("province"):
        queryset = queryset.filter(province=filters["province"])
    if filters.get("district"):
        queryset = queryset.filter(district=filters["district"])
    return queryset


def _apply_feedback_filters(queryset, filters):
    q = filters.get("q")
    if q:
        queryset = queryset.filter(
            Q(contact_name__icontains=q)
            | Q(contact_email__icontains=q)
            | Q(role_or_group__icontains=q)
            | Q(page_url__icontains=q)
            | Q(comment__icontains=q)
            | Q(status__icontains=q)
            | Q(user__username__icontains=q)
            | Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q)
        )
    return queryset


def _yes_no(value):
    return "Yes" if value else "No"


def _maybe_slice(rows, limit):
    rows = list(rows)
    if limit is None:
        return rows
    return rows[:limit]


def build_report(request, filters, limit=None):
    report_type = filters.get("report") or "workers"
    if report_type == "disability":
        return build_disability_report(request, filters, limit)
    if report_type == "workplaces":
        return build_workplace_report(request, filters, limit)
    if report_type == "labour":
        return build_labour_report(request, filters, limit)
    if report_type == "activities":
        return build_activity_report(request, filters, limit)
    if report_type == "feedback":
        return build_feedback_report(request, filters, limit)
    return build_worker_report(request, filters, limit)


def build_worker_report(request, filters, limit=None):
    queryset = _apply_worker_filters(_worker_queryset(request), filters)
    headers = [
        "Unique ID",
        "Full name",
        "Gender",
        "Age",
        "Phone",
        "Province",
        "District",
        "Workplace",
        "Sector",
        "Occupation",
        "Employment status",
        "Union membership",
        "Union",
        "Disability type",
        "Severity",
        "Captured by",
        "Capture source",
        "Capture mode",
        "Device captured",
        "Verified",
    ]
    rows = (
        [
            worker.unique_id,
            worker.full_name,
            worker.get_gender_display(),
            worker.age or "",
            worker.phone,
            worker.province,
            worker.district,
            worker.workplace.employer_name if worker.workplace else "",
            worker.get_sector_display(),
            worker.occupation,
            worker.get_employment_status_display(),
            worker.get_union_membership_status_display(),
            worker.union_name,
            worker.disability_profile.get_disability_type_display() if hasattr(worker, "disability_profile") else "",
            worker.disability_profile.get_severity_display() if hasattr(worker, "disability_profile") else "",
            worker.captured_by_display,
            worker.get_capture_source_display(),
            worker.get_capture_mode_display(),
            worker.captured_at_device.strftime("%Y-%m-%d %H:%M") if worker.captured_at_device else "",
            _yes_no(worker.data_verified),
        ]
        for worker in queryset
    )
    return report_payload("workers", "Worker Register", queryset.count(), headers, _maybe_slice(rows, limit), limit)


def build_labour_report(request, filters, limit=None):
    queryset = _apply_worker_filters(_worker_queryset(request), filters)
    headers = [
        "Unique ID",
        "Worker",
        "Province",
        "District",
        "Union",
        "Grievance reported",
        "Case status",
        "Legal support",
        "Collective bargaining",
        "Social protection",
        "Captured by",
        "Capture source",
        "Capture mode",
        "Grievance summary",
    ]
    rows = []
    for worker in queryset:
        labour = getattr(worker, "labour_rights", None)
        rows.append(
            [
                worker.unique_id,
                worker.full_name,
                worker.province,
                worker.district,
                worker.union_name,
                _yes_no(labour and labour.grievance_reported),
                labour.get_case_status_display() if labour else "No active case",
                _yes_no(labour and labour.legal_support_provided),
                _yes_no(labour and labour.collective_bargaining_coverage),
                _yes_no(labour and labour.social_protection_coverage),
                worker.captured_by_display,
                worker.get_capture_source_display(),
                worker.get_capture_mode_display(),
                labour.grievance_summary if labour else "",
            ]
        )
    return report_payload("labour", "Labour Rights Report", queryset.count(), headers, _maybe_slice(rows, limit), limit)


def build_disability_report(request, filters, limit=None):
    queryset = _apply_worker_filters(_worker_queryset(request), filters)
    disability_labels = dict(DISABILITY_TYPE_CHOICES)
    severity_labels = dict(SEVERITY_CHOICES)
    headers = ["Province", "District", "Disability type", "Severity", "Workers"]
    grouped = (
        queryset.values("province", "district", "disability_profile__disability_type", "disability_profile__severity")
        .annotate(total=Count("id"))
        .order_by("province", "district", "disability_profile__disability_type", "disability_profile__severity")
    )
    rows = [
        [
            row["province"],
            row["district"],
            disability_labels.get(row["disability_profile__disability_type"], "Not recorded"),
            severity_labels.get(row["disability_profile__severity"], "Not recorded"),
            row["total"],
        ]
        for row in grouped
    ]
    return report_payload("disability", "Disability Summary", len(rows), headers, _maybe_slice(rows, limit), limit)


def build_workplace_report(request, filters, limit=None):
    queryset = _apply_workplace_filters(_workplace_queryset(request), filters)
    headers = [
        "Employer",
        "Workplace",
        "Province",
        "District",
        "Sector",
        "Accessibility status",
        "Accommodations available",
        "Union presence",
        "Workers",
    ]
    rows = (
        [
            workplace.employer_name,
            workplace.workplace_name,
            workplace.province,
            workplace.district,
            workplace.get_sector_display(),
            workplace.get_accessibility_status_display(),
            _yes_no(workplace.accommodations_available),
            _yes_no(workplace.union_presence),
            getattr(workplace, "worker_count", workplace.workers.count()),
        ]
        for workplace in queryset
    )
    return report_payload("workplaces", "Workplace Accessibility Report", queryset.count(), headers, _maybe_slice(rows, limit), limit)


def build_activity_report(request, filters, limit=None):
    queryset = _apply_activity_filters(ProjectActivity.objects.all(), filters)
    headers = [
        "Activity",
        "Type",
        "Partner",
        "Province",
        "District",
        "Start date",
        "Participants",
        "Participants with disabilities",
    ]
    rows = (
        [
            activity.title,
            activity.get_activity_type_display(),
            activity.partner,
            activity.province,
            activity.district,
            activity.start_date,
            activity.participants_total,
            activity.participants_with_disabilities,
        ]
        for activity in queryset
    )
    return report_payload("activities", "Project Monitoring Activities", queryset.count(), headers, _maybe_slice(rows, limit), limit)


def build_feedback_report(request, filters, limit=None):
    queryset = _apply_feedback_filters(_feedback_queryset(request), filters)
    headers = [
        "Submitted",
        "Contact name",
        "Contact email",
        "Role / group",
        "Submitted by",
        "Page",
        "Status",
        "Comment",
    ]
    rows = (
        [
            feedback.created_at.strftime("%Y-%m-%d %H:%M"),
            feedback.contact_name,
            feedback.contact_email,
            feedback.role_or_group,
            feedback.user.get_full_name() or feedback.user.username if feedback.user else "",
            feedback.page_url,
            feedback.get_status_display(),
            feedback.comment,
        ]
        for feedback in queryset
    )
    return report_payload("feedback", "Stakeholder Feedback Report", queryset.count(), headers, _maybe_slice(rows, limit), limit)


def report_payload(report_type, title, total_count, headers, rows, limit):
    return {
        "type": report_type,
        "title": title,
        "total_count": total_count,
        "headers": headers,
        "rows": rows,
        "is_limited": limit is not None and total_count > limit,
        "limit": limit,
    }


def _validated_filters(request):
    form = ReportFilterForm(request.GET or {"report": "workers"})
    if form.is_valid():
        return form.cleaned_data
    fallback = ReportFilterForm({"report": "workers"})
    fallback.is_valid()
    return fallback.cleaned_data


@login_required
def export_workers_csv(request):
    report = build_report(request, _validated_filters(request), limit=None)
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = f'attachment; filename="dldms-{report["type"]}.csv"'
    writer = csv.writer(response)
    writer.writerow(report["headers"])
    writer.writerows(report["rows"])
    return response


@login_required
def export_workers_xlsx(request):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    report = build_report(request, _validated_filters(request), limit=None)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report["title"][:31]
    sheet.append(report["headers"])
    for row in report["rows"]:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 42)
        sheet.column_dimensions[column_cells[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="dldms-{report["type"]}.xlsx"'
    return response


@login_required
def summary_pdf(request):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    report = build_report(request, _validated_filters(request), limit=None)
    output = BytesIO()
    page_size = landscape(A4)
    document = SimpleDocTemplate(output, pagesize=page_size, title=report["title"])
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Disability-Disaggregated Labour Data Management System", styles["Title"]),
        Paragraph(report["title"], styles["Heading2"]),
        Paragraph(f'Total rows: {report["total_count"]}', styles["Normal"]),
        Spacer(1, 12),
    ]

    table_data = [report["headers"]] + [
        [truncate_cell(value) for value in row]
        for row in report["rows"]
    ]
    usable_width = page_size[0] - document.leftMargin - document.rightMargin
    col_width = usable_width / max(len(report["headers"]), 1)
    table = Table(table_data or [["No data"]], colWidths=[col_width] * len(report["headers"]), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 6),
                ("PADDING", (0, 0), (-1, -1), 3),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    document.build(story)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="dldms-{report["type"]}.pdf"'
    return response


def truncate_cell(value, limit=42):
    text = str(value or "")
    return text if len(text) <= limit else f"{text[: limit - 3]}..."
